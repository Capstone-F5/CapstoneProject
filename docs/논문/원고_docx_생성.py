# -*- coding: utf-8 -*-
"""논문초안_본문.md -> 학회 양식 서식을 입힌 워드 원고 생성기.

출력 (양식 파일은 읽기만 하며 절대 수정하지 않는다)
  논문_원고.docx        저자 정보 포함 (최종논문 제출용)
  논문_원고_심사용.docx  저자 정보·저자소개 제외 (심사논문 제출용, 양식 지침)

생성된 파일을 양식 파일에 직접 반영하는 것은 사용자가 수동으로 한다.
양식의 문단 서식(휴먼명조/Times New Roman, 1.5줄, 장·절 제목 정렬, 참고문헌 자동번호)을
그대로 복제하며, 마크다운의 *...* 는 이탤릭 run으로 변환한다.
"""
import copy, io, os, re, shutil, sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '한국실천공학교육학회_논문지_양식.docx')
MD = os.path.join(HERE, '논문초안_본문.md')

# ---------------------------------------------------------------- 원고 파싱
md = open(MD, encoding='utf-8').read()


def section(a, b):
    return md.split(a)[1].split(b)[0]


def paras(text):
    out = []
    for blk in text.strip().split('\n\n'):
        blk = blk.strip()
        if blk and not blk.startswith(('>', '#', '|', '-', '**Key Words')):
            out.append(' '.join(blk.split()))
    return out


def author(label):
    # 마크다운에서 각주 기호(*)가 강조로 읽히지 않도록 백틱으로 감싸 두었다
    return re.search(r'^- %s:\s*(.+)$' % label, md, re.M).group(1).strip().strip('`')


AUTHORS = {k: author(k) for k in ('국문 성명', '국문 소속', '영문 성명', '영문 소속')}
summary = paras(section('## 요 약', '## ABSTRACT'))[0]
abstract = paras(section('## ABSTRACT', '**Key Words:**'))[0]
keywords = re.search(r'\*\*Key Words:\*\*\s*(.+)', md).group(1).strip()
intro = paras(section('## I. 서론', '## II.'))

ch2_title = re.search(r'^## (II\..+)$', md, re.M).group(1).strip()
ch2 = []
for blk in section('## %s' % ch2_title, '## References').strip().split('\n\n'):
    blk = blk.strip()
    if blk.startswith('### '):
        ch2.append(('sec', blk[4:].strip()))
    elif blk:
        ch2.append(('body', ' '.join(blk.split())))

refs = [re.sub(r'^\[\d+\]\s*', '', l).strip()
        for l in section('## References', '## 그림 및 표').splitlines()
        if re.match(r'^\[\d+\]', l)]

fig_tbl = md.split('## 그림 및 표')[1].split('\n---\n')[0]
tbl_rows = [[c.strip() for c in l.strip().strip('|').split('|')]
            for l in fig_tbl.splitlines()
            if l.strip().startswith('|') and not re.match(r'^\|[\s\-|]+\|$', l.strip())]
assert len(tbl_rows) >= 5 and len(tbl_rows[0]) == 6, '표 1 파싱 실패'

SKELETON = [
    ('chap', 'III. 제안 시스템'),
    ('sec', 'A. 설계 요구사항과 전체 구조'), ('body', '(작성 예정)'),
    ('sec', 'B. 멀티모달 상호작용 모델과 모드 전환'), ('body', '(작성 예정)'),
    ('sec', 'C. LLM 음성 주문 파이프라인과 근거화된 실행'), ('body', '(작성 예정)'),
    ('sec', 'D. 손동작 인식과 보조기기 감지'), ('body', '(작성 예정)'),
    ('sec', 'E. 구현 환경'), ('body', '(작성 예정)'),
    ('chap', 'IV. 평가 프레임워크'),
    ('sec', 'A. 평가 설계 원칙'), ('body', '(작성 예정)'),
    ('sec', 'B. 3계층 평가 지표'), ('body', '(작성 예정)'),
    ('sec', 'C. 연구 윤리 및 개인정보'), ('body', '(작성 예정)'),
    ('chap', 'V. 결 론'), ('body', '(작성 예정)'),
    ('chap', '감사의 글'), ('body', '(작성 예정)'),
]


LATIN, HANGUL = 'Times New Roman', '휴먼명조'


def enforce_fonts(doc):
    """모든 run과 문단 기호에 영문/한글 글꼴을 함께 지정한다 (양식: 휴먼명조, 영문 Times New Roman)."""
    def apply(rpr):
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for attr, val in (('w:ascii', LATIN), ('w:hAnsi', LATIN),
                          ('w:cs', LATIN), ('w:eastAsia', HANGUL)):
            rf.set(qn(attr), val)

    n = 0
    for r in doc.element.body.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            r.insert(0, rpr)
        apply(rpr)
        n += 1
    for ppr in doc.element.body.iter(qn('w:pPr')):   # 문단 기호 서식
        rpr = ppr.find(qn('w:rPr'))
        if rpr is not None:
            apply(rpr)
    # 스타일 기본값도 맞춰 두어 새로 입력하는 글자에도 적용되게 한다
    normal = doc.styles['Normal'].element
    rpr = normal.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        normal.append(rpr)
    apply(rpr)
    return n


# ---------------------------------------------------------------- 문서 생성
def build(out_path, include_authors):
    shutil.copyfile(SRC, out_path)          # 양식은 읽기 전용으로만 사용
    doc = Document(out_path)

    def find(pred):
        for p in doc.paragraphs:
            if pred(p.text):
                return p
        raise KeyError('문단을 찾지 못함')

    def set_text(p, text, rich=False):
        """문단 서식은 유지한 채 본문만 교체한다.

        rich=True 일 때만 *...* 를 이탤릭 run으로 바꾼다. 저자 각주 기호(*, **)가
        이탤릭 markup 으로 오인되지 않도록 참고문헌에만 적용한다.
        """
        rpr = None
        if p.runs and p.runs[0]._r.rPr is not None:
            rpr = copy.deepcopy(p.runs[0]._r.rPr)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        chunks = re.split(r'\*([^*\n]+)\*', text) if rich else [text]
        for i, chunk in enumerate(chunks):
            if not chunk:
                continue
            run = p.add_run(chunk)
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
            if rich and i % 2 == 1:         # re.split의 홀수 인덱스 = *...* 안쪽
                run.italic = True
        if not p.runs:
            p.add_run('')
        return p

    def drop(p):
        p._p.getparent().remove(p._p)

    TPL = {k: copy.deepcopy(find(lambda t, s=s: t.startswith(s))._p) for k, s in
           (('body', '디지털 전환과 비대면'), ('chap', 'I. 서론 - 수준 1'),
            ('sec', 'A. 수준 2 제목'), ('ref', 'E. J. Sin and S. B. Lim'),
            ('cap', '그림 1. 한국실천공학교육학회'))}

    class Cursor:
        def __init__(self, anchor):
            self.node = anchor._p

        def add(self, kind, text):
            new = copy.deepcopy(TPL[kind])
            self.node.addnext(new)
            self.node = new
            set_text(Paragraph(new, None), text, rich=(kind == 'ref'))
            return self

    # 1) 상단 유의사항 표 제거
    t0 = doc.tables[0]._tbl
    t0.getparent().remove(t0)

    # 2) 저자 정보: 최종본은 채우고, 심사용은 통째로 제거
    drop(find(lambda t: t.startswith('아래의 성명, 소속 및 직위는')))
    slots = [('성 춘 향*, 김 길 동**', '국문 성명'),
             ('*한국대학교 교육공학과', '국문 소속'),
             ('Chun-hyang Sung*, Gil-dong Kim**', '영문 성명'),
             ('*Professor, Department of Education', '영문 소속')]
    for key, field in slots:
        p = find(lambda t, k=key: t.startswith(k))
        set_text(p, AUTHORS[field]) if include_authors else drop(p)

    # 3) 요약 / ABSTRACT / Key Words
    set_text(find(lambda t: t.startswith('요 약 (10포인트')), '요 약')
    set_text(find(lambda t: t.startswith('이곳에 요약문을 작성합니다')), summary)
    set_text(find(lambda t: t.startswith('ABSTRACT (10포인트')), 'ABSTRACT')
    set_text(find(lambda t: t.startswith('Please write the abstract')), abstract)
    set_text(find(lambda t: t.startswith('Key Words : 본문 중에')), 'Key Words : ' + keywords)
    drop(find(lambda t: t.startswith('예) Key Words')))

    # 4) 'I. 서론' 제목만 남기고 그 뒤 ~ '참 고 문 헌' 앞까지 양식 예시를 비운다
    intro_head = set_text(find(lambda t: t.startswith('I. 서론 - 수준 1')), 'I. 서론')
    refs_head = find(lambda t: t.startswith('참 고 문 헌'))._p
    node = intro_head._p.getnext()
    while node is not None and node is not refs_head:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt

    # 5) 서론 ~ V장 본문
    cur = Cursor(intro_head)
    for txt in intro:
        cur.add('body', txt)
    cur.add('chap', ch2_title)
    for kind, text in ch2:
        cur.add(kind, text)
    for kind, text in SKELETON:
        cur.add(kind, text)

    # 6) 참고문헌 (양식의 자동 번호 목록을 그대로 이어 쓴다)
    ref_ps = [p for p in doc.paragraphs if p.style.name == 'List Paragraph']
    for i, p in enumerate(ref_ps):
        if i < len(refs):
            set_text(p, refs[i], rich=True)
        else:
            drop(p)
    cur = Cursor(ref_ps[min(len(refs), len(ref_ps)) - 1])
    for text in refs[len(ref_ps):]:
        cur.add('ref', text)

    # 7) 그림·표 안내문과 로고 예시 제거, 표 1 캡션으로 교체
    for key in ('모든 그림이나 표는 본문 내에서', '표의 사용도 그림과 동일한 형식을',
                '참고문헌 들은 반드시 본문 내에서', '참고문헌은 반드시 영문으로 작성하며'):
        drop(find(lambda t, k=key: t.startswith(k)))
    drop(find(lambda t: t.strip() == '그림과 표'))
    logo = find(lambda t: t.startswith('그림 1. 한국실천공학교육학회'))
    prev = logo._p.getprevious()
    if prev is not None and not Paragraph(prev, None).text.strip():
        prev.getparent().remove(prev)
    cap_ko = set_text(logo, '표 1. 선행연구와 제안 시스템의 비교')
    cap_en = set_text(find(lambda t: t.startswith('Fig. 1. KIPEE logo')),
                      'Table 1. Comparison of related work and the proposed system')
    for c in (cap_ko, cap_en):
        c.alignment = None
    drop(find(lambda t: t.startswith('표 1. 시뮬레이션 파라미터')))
    drop(find(lambda t: t.startswith('Table 1. Simulation parameters')))
    node = cap_en._p.getnext()
    while node is not None and node.tag.endswith('}p') and not Paragraph(node, None).text.strip():
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt

    # 8) 예시 표 -> 표 1
    old = doc.tables[-1]
    new = doc.add_table(rows=len(tbl_rows), cols=len(tbl_rows[0]))
    new.style = 'Table Grid'
    for r, row in enumerate(tbl_rows):
        for c, val in enumerate(row):
            set_text(new.cell(r, c).paragraphs[0], val)
            if r == 0:
                new.cell(r, c).paragraphs[0].runs[0].bold = True
    old._tbl.addnext(new._tbl)
    old._tbl.getparent().remove(old._tbl)

    # 9) 저자소개 예시 제거 (약력·사진은 최종 제출 시 수동 추가)
    node = find(lambda t: t.startswith('저 자 소 개'))._p
    while node is not None:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt

    # 10) 글꼴 통일: 영문 Times New Roman, 한글 휴먼명조
    #     워드는 run 안에서 문자 종류별로 ascii / eastAsia 글꼴을 따로 적용하므로
    #     두 값을 모두 지정하면 한 문단에 섞여 있어도 각각 올바른 글꼴로 나온다.
    enforce_fonts(doc)

    doc.save(out_path)
    return doc


print('파싱: 요약 %d자 / ABSTRACT %d단어 / 서론 %d문단 / %s %d블록 / 참고문헌 %d편 / 표 %dx%d'
      % (len(summary), len(abstract.split()), len(intro), ch2_title.split('.')[0],
         len(ch2), len(refs), len(tbl_rows), len(tbl_rows[0])))
for name, with_authors in (('논문_원고.docx', True), ('논문_원고_심사용.docx', False)):
    build(os.path.join(HERE, name), with_authors)
    print('  생성: %-22s 저자 정보 %s' % (name, '포함' if with_authors else '제외'))
